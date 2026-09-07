import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from ...config import settings
from .generator import classify_table_headers

DEFAULT_CHUNK_MAX_CHARS = settings.chunking.chunk_max_chars

HEADING_STYLES = {"heading 1", "heading 2", "heading 3", "title"}

QUESTION_HEADERS = {"question", "questions", "question(s)", "query"}
ANSWER_HEADERS = {"answer", "answers", "response", "responses", "solution", "solutions"}
CATEGORY_HEADERS = {"category"}
SUBCATEGORY_HEADERS = {"sub category", "subcategory"}
SERIAL_HEADERS = {"serial no", "serial number", "serial", "sno", "s no", "sl no", "srno", "sr no", "sr", "#"}


@dataclass
class Chunk:
    section: str
    text: str


@dataclass
class TableQA:
    question: str
    answer: str
    category: str = ""
    sub_category: str = ""
    serial_no: str = ""


@dataclass
class Extraction:
    chunks: list[Chunk] = field(default_factory=list)
    table_qa: list[TableQA] = field(default_factory=list)


def _normalize_header(text: str) -> str:
    return re.sub(r"[\s_.\-]+", " ", text.strip().lower()).strip()


def _find_col(header: list[str], names: set[str]) -> int | None:
    for i, cell in enumerate(header):
        if _normalize_header(cell) in names:
            return i
    return None


def _parse_qa_table(header: list[str], rows: list[list[str]], chat_model=None) -> list[TableQA] | None:
    """If header looks like a Question/Answer table, extract rows directly; else None."""
    q_idx = _find_col(header, QUESTION_HEADERS)
    a_idx = _find_col(header, ANSWER_HEADERS)
    cat_idx = _find_col(header, CATEGORY_HEADERS)
    sub_idx = _find_col(header, SUBCATEGORY_HEADERS)
    ser_idx = _find_col(header, SERIAL_HEADERS)

    if (q_idx is None or a_idx is None) and chat_model is not None:
        try:
            mapped = classify_table_headers(chat_model, header)
        except ValueError:
            mapped = {}
        q_idx = q_idx if q_idx is not None else mapped.get("question_col")
        a_idx = a_idx if a_idx is not None else mapped.get("answer_col")
        cat_idx = cat_idx if cat_idx is not None else mapped.get("category_col")
        sub_idx = sub_idx if sub_idx is not None else mapped.get("sub_category_col")
        ser_idx = ser_idx if ser_idx is not None else mapped.get("serial_col")

    if q_idx is None or a_idx is None:
        return None

    qa_rows = []
    for row in rows:
        question = row[q_idx].strip() if q_idx < len(row) else ""
        answer = row[a_idx].strip() if a_idx < len(row) else ""
        if not question or not answer:
            continue
        category = row[cat_idx].strip() if cat_idx is not None and cat_idx < len(row) else ""
        sub_category = row[sub_idx].strip() if sub_idx is not None and sub_idx < len(row) else ""
        serial_no = row[ser_idx].strip() if ser_idx is not None and ser_idx < len(row) else ""
        qa_rows.append(
            TableQA(
                question=question,
                answer=answer,
                category=category,
                sub_category=sub_category,
                serial_no=serial_no,
            )
        )

    return qa_rows or None


def extract_docx(path: Path, chunk_max_chars: int = DEFAULT_CHUNK_MAX_CHARS, chat_model=None) -> Extraction:
    doc = Document(str(path))
    result = Extraction()
    section = "Document"
    buffer: list[str] = []

    def flush():
        text = "\n".join(buffer).strip()
        if text:
            result.chunks.append(Chunk(section=section, text=text))
        buffer.clear()

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style in HEADING_STYLES:
            flush()
            section = text
            continue
        buffer.append(text)
        if sum(len(b) for b in buffer) > chunk_max_chars:
            flush()

    flush()

    for table in doc.tables:
        table_rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not table_rows:
            continue

        header, data_rows = table_rows[0], table_rows[1:]
        qa_rows = _parse_qa_table(header, data_rows, chat_model) if data_rows else None
        if qa_rows:
            result.table_qa.extend(qa_rows)
            continue

        table_text = "\n".join(" | ".join(r) for r in table_rows).strip()
        if table_text:
            result.chunks.append(Chunk(section=f"{section} (table)", text=table_text))

    return result


def _rows_to_text(header: list[str], rows: list[list[str]]) -> str:
    lines = [" | ".join(header)] if header else []
    for row in rows:
        lines.append(" | ".join(row))
    return "\n".join(lines).strip()


def extract_xlsx(path: Path, chunk_max_chars: int = DEFAULT_CHUNK_MAX_CHARS, chat_model=None) -> Extraction:
    wb = load_workbook(str(path), data_only=True, read_only=True)
    result = Extraction()

    for sheet in wb.worksheets:
        all_rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                all_rows.append(cells)

        if not all_rows:
            continue

        header, data_rows = all_rows[0], all_rows[1:]

        qa_rows = _parse_qa_table(header, data_rows, chat_model) if data_rows else None
        if qa_rows:
            result.table_qa.extend(qa_rows)
            continue

        if not data_rows:
            text = _rows_to_text(header, [])
            if text:
                result.chunks.append(Chunk(section=sheet.title, text=text))
            continue

        batch: list[list[str]] = []
        batch_chars = 0
        part = 1
        for row in data_rows:
            batch.append(row)
            batch_chars += sum(len(c) for c in row)
            if batch_chars > chunk_max_chars:
                label = sheet.title if part == 1 else f"{sheet.title} (part {part})"
                result.chunks.append(Chunk(section=label, text=_rows_to_text(header, batch)))
                batch, batch_chars = [], 0
                part += 1

        if batch:
            label = sheet.title if part == 1 else f"{sheet.title} (part {part})"
            result.chunks.append(Chunk(section=label, text=_rows_to_text(header, batch)))

    wb.close()
    return result


def extract_pdf(path: Path, chunk_max_chars: int = DEFAULT_CHUNK_MAX_CHARS) -> Extraction:
    reader = PdfReader(str(path))
    result = Extraction()
    buffer: list[str] = []
    buffer_chars = 0
    start_page = 1

    def flush(end_page: int):
        nonlocal buffer, buffer_chars, start_page
        text = "\n".join(buffer).strip()
        if text:
            label = f"Page {start_page}" if start_page == end_page else f"Page {start_page}-{end_page}"
            result.chunks.append(Chunk(section=label, text=text))
        buffer, buffer_chars = [], 0
        start_page = end_page + 1

    for page_num, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        buffer.append(text)
        buffer_chars += len(text)
        if buffer_chars > chunk_max_chars:
            flush(page_num)

    flush(len(reader.pages))
    return result


def extract_document(path: Path, chunk_max_chars: int = DEFAULT_CHUNK_MAX_CHARS, chat_model=None) -> Extraction:
    ext = path.suffix.lower()
    if ext == ".docx":
        return extract_docx(path, chunk_max_chars, chat_model)
    if ext == ".xlsx":
        return extract_xlsx(path, chunk_max_chars, chat_model)
    if ext == ".pdf":
        return extract_pdf(path, chunk_max_chars)
    raise ValueError(f"Unsupported file type: {ext}")
